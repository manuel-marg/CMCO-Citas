"""
Extracción de la hoja 'Base de Datos' desde el Excel oficial hacia database/pacientes.sql
Centro Médico Comunitario OscaryAnny (CMCO / ADH)
"""
import openpyxl
import re
import os

def clean_str(val):
    if val is None:
        return None
    s = str(val).strip()
    if s == '' or s.lower() in ['none', 'null', '0', 's/c']:
        return None
    return s

def escape_sql(val):
    if val is None:
        return 'NULL'
    s = str(val).replace('\\', '\\\\').replace("'", "\\'")
    return f"'{s}'"

def main():
    excel_file = 'Centro Médico Comunitario OscaryAnny - CONTROL DE CITAS Sistema.xlsx'
    if not os.path.exists(excel_file):
        print(f"Error: No se encontró {excel_file}")
        return

    print(f"Cargando {excel_file}...")
    wb = openpyxl.load_workbook(excel_file, read_only=True)
    if 'Base de Datos' not in wb.sheetnames:
        print("Error: No existe la hoja 'Base de Datos' en el libro Excel.")
        return

    ws = wb['Base de Datos']
    
    pacientes_map = {}
    total_filas = 0

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue  # Saltar encabezados
            
        nombre_raw = clean_str(row[0])
        if not nombre_raw:
            continue
            
        total_filas += 1
        
        repres_raw = clean_str(row[1])
        
        # Procesar Teléfono
        tel_val = row[2]
        telefono = None
        if tel_val is not None:
            if isinstance(tel_val, float):
                tel_val = int(tel_val)
            t_str = str(tel_val).strip()
            t_digits = re.sub(r'\D', '', t_str)
            if len(t_digits) == 10 and not t_digits.startswith('0'):
                telefono = '0' + t_digits
            elif len(t_digits) >= 7:
                telefono = t_digits
            elif t_str.lower() not in ['', 'none', 'null', '0']:
                telefono = t_str
                
        # Procesar Cédula
        ced_val = row[3]
        cedula = None
        if ced_val is not None:
            c_str = str(ced_val).strip()
            if c_str.lower() not in ['', 'none', 'null', '0', 's/c']:
                cedula = c_str

        # Clasificación de Tipo de Beneficio según Columna 'Empleado'
        rol_raw = clean_str(row[4])
        tipo_beneficio = 'Particular'
        observaciones = None
        
        if rol_raw:
            r_lower = rol_raw.lower()
            if any(w in r_lower for w in ['emplead', 'vice-presidente', 'vice-presidenta', 'presidente', 'personal']):
                tipo_beneficio = 'Empleado'
            elif any(w in r_lower for w in ['familiar', 'hijo', 'hija']):
                tipo_beneficio = 'Familiar'
            elif 'participante' in r_lower:
                tipo_beneficio = 'Participante'
            else:
                tipo_beneficio = rol_raw

            if rol_raw.lower() != tipo_beneficio.lower():
                observaciones = f"Rol original: {rol_raw}"

        # Clave única para consolidar registros repetidos en el Excel sin perder datos
        key = (nombre_raw.lower(), cedula.lower() if cedula else '')
        if key in pacientes_map:
            p_exist = pacientes_map[key]
            if not p_exist['telefono'] and telefono:
                p_exist['telefono'] = telefono
            if not p_exist['representante'] and repres_raw:
                p_exist['representante'] = repres_raw
            if not p_exist['cedula'] and cedula:
                p_exist['cedula'] = cedula
            if p_exist['tipo_beneficio'] == 'Particular' and tipo_beneficio != 'Particular':
                p_exist['tipo_beneficio'] = tipo_beneficio
            if not p_exist['observaciones'] and observaciones:
                p_exist['observaciones'] = observaciones
        else:
            pacientes_map[key] = {
                'nombre': nombre_raw,
                'cedula': cedula,
                'telefono': telefono,
                'representante': repres_raw,
                'tipo_beneficio': tipo_beneficio,
                'observaciones': observaciones
            }

    pacientes_lista = list(pacientes_map.values())
    pacientes_lista.sort(key=lambda p: p['nombre'].lower())

    print(f"Filas leídas en Excel: {total_filas}")
    print(f"Pacientes únicos consolidados: {len(pacientes_lista)}")

    # Construir contenido del archivo SQL
    sql_lines = [
        "-- =============================================================================",
        "-- CMCO-Citas - Catálogo Completo de Pacientes",
        "-- Centro Médico Comunitario OscaryAnny / Asociación Dios Habla (ADH)",
        f"-- Total Registros: {len(pacientes_lista)} pacientes consolidados",
        "-- =============================================================================",
        "",
        "SET NAMES utf8mb4 COLLATE utf8mb4_unicode_ci;",
        "USE `cmco_citas_db`;",
        "SET FOREIGN_KEY_CHECKS = 0;",
        "",
        "-- Inserción de pacientes",
    ]

    for p in pacientes_lista:
        line = (
            f"INSERT INTO `pacientes` (`nombre`, `cedula`, `telefono`, `representante`, `tipo_beneficio`, `observaciones`) "
            f"VALUES ({escape_sql(p['nombre'])}, {escape_sql(p['cedula'])}, {escape_sql(p['telefono'])}, "
            f"{escape_sql(p['representante'])}, {escape_sql(p['tipo_beneficio'])}, {escape_sql(p['observaciones'])});"
        )
        sql_lines.append(line)

    sql_lines.append("")
    sql_lines.append("SET FOREIGN_KEY_CHECKS = 1;")
    sql_lines.append("")

    out_path = os.path.join('database', 'pacientes.sql')
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(sql_lines))

    print(f"Archivo generado exitosamente en: {out_path} ({os.path.getsize(out_path)} bytes)")

if __name__ == '__main__':
    main()
